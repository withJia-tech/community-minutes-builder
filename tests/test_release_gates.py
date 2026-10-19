"""Release-gate regressions: these must pass before considering a v0.1.0 tag."""
import json
import os
from pathlib import Path
from types import SimpleNamespace
import zipfile

import pytest
from lxml import etree
from PIL import Image
from docx import Document
from docx.shared import Inches

import build
from helpers import FIXTURES, lines, run_script, write_json

PROFILE = FIXTURES / 'committee-profile.json'
SOURCES = FIXTURES / 'row-sources.json'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def _remove_tables(source, target):
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(target, 'w') as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == 'word/document.xml':
                root = etree.fromstring(data)
                body = root.find(W + 'body')
                for table in body.findall(W + 'tbl'):
                    body.remove(table)
                data = etree.tostring(root)
            dst.writestr(item, data)
    return target


def test_empty_document_fails_both_linter_and_row_audit(tmp_path):
    source = build.build_docx(tmp_path / 'source.docx')
    empty = _remove_tables(source, tmp_path / 'empty.docx')
    lint = run_script('lint_minutes.py', empty, '--final')
    audit = run_script('row_audit.py', empty, '--source-map', SOURCES)
    assert lint.returncode == 1 and 'expected an attendance table' in lint.stdout
    assert audit.returncode == 1 and 'OVERALL VERDICT: FAIL' in audit.stdout
    assert 'Traceback' not in audit.stderr


@pytest.mark.parametrize('data', [b'not a zip', b'', b'PK\x03\x04unreadable'])
def test_malformed_docx_fails_cleanly(tmp_path, data):
    file = tmp_path / 'bad.docx'
    file.write_bytes(data)
    for cmd in ('lint_minutes.py', 'row_audit.py'):
        r = run_script(cmd, file)
        assert r.returncode == 1, (cmd, r.stdout, r.stderr)
        assert 'Traceback' not in r.stderr


def test_duplicate_numbers_and_unnumbered_content_are_rejected(tmp_path):
    items = list(build.DRAFT_ITEMS)
    items[2] = ('2', items[2][1], items[2][2])
    duplicate = build.build_docx(tmp_path / 'duplicate.docx', items=items)
    missing = build.build_docx(tmp_path / 'missing.docx', items=[('', 'A silent issue', '')])
    for file, diagnostic in ((duplicate, 'duplicate minutes row number 2'),
                             (missing, 'invalid/missing number')):
        for cmd in ('lint_minutes.py', 'row_audit.py'):
            r = run_script(cmd, file)
            assert r.returncode == 1 and diagnostic in r.stdout, r.stdout


def test_first_body_title_missing_is_blocker_even_if_roster_has_titles(tmp_path):
    items = list(build.DRAFT_ITEMS)
    items[0] = ('1', 'Confirmation. Alice Lim presented the agenda.', '')
    items[2] = ('3', "Chairman's Report. Mdm Alice Lim spoke.", '')
    file = build.build_docx(tmp_path / 'no-title.docx', items=items)
    final = run_script('lint_minutes.py', file, '--profile', PROFILE, '--final')
    assert final.returncode == 1
    assert "'Alice Lim' first appears in minutes row 1 without an approved title" in final.stdout
    fixed_items = list(items)
    fixed_items[0] = ('1', 'Confirmation. Chairman Alice Lim presented the agenda.', '')
    fixed = build.build_docx(tmp_path / 'fixed.docx', items=fixed_items)
    assert run_script('lint_minutes.py', fixed, '--profile', PROFILE, '--final').returncode == 0


def test_split_runs_are_reconstructed_for_placeholders_and_titles(tmp_path):
    """python-docx generates a Word-compatible package with split runs."""
    doc = _make_word_compatible_doc()
    file = tmp_path / 'split.docx'
    row = doc.tables[-1].rows[0].cells[1]
    row.text = ''
    row.paragraphs[0].add_run('Chairman Alice Lim said [TB')
    row.paragraphs[0].add_run('C]')
    doc.save(file)
    result = run_script('lint_minutes.py', file, '--profile', PROFILE, '--final')
    assert result.returncode == 1
    assert "matching '[TBC]'" in result.stdout
    assert 'without an approved title' not in result.stdout


def _make_word_compatible_doc():
    doc = Document()
    attendance = doc.add_table(rows=0, cols=4)
    row = attendance.add_row()
    row.cells[0].merge(row.cells[3]).text = 'Present'
    attendee = attendance.add_row()
    attendee.cells[0].text = '1'
    attendee.cells[1].text = 'Mdm'
    attendee.cells[3].text = 'Alice Lim'
    row = attendance.add_row()
    row.cells[0].merge(row.cells[3]).text = 'Absent with Apologies'
    row = attendance.add_row()
    row.cells[0].merge(row.cells[3]).text = 'In Attendance'
    table = doc.add_table(rows=1, cols=3)
    table.cell(0, 0).text = '1'
    table.cell(0, 1).text = 'Chairman Alice Lim presented the report.'
    table.cell(0, 2).text = ''
    return doc


def test_python_docx_generated_table_passes(tmp_path):
    """Word-compatible fixture, not a claim of testing actual MS Word output."""
    document = _make_word_compatible_doc()
    file = tmp_path / 'created-with-python-docx.docx'
    document.save(file)
    lint = run_script('lint_minutes.py', file, '--profile', PROFILE, '--final')
    assert lint.returncode == 0, lint.stdout
    sources = write_json(tmp_path / 'sources.json', {'1': 'Reviewed secretary notes'})
    assert run_script('row_audit.py', file, '--source-map', sources).returncode == 0


def test_nested_table_does_not_change_detected_minutes_table(tmp_path):
    document = _make_word_compatible_doc()
    nested = document.tables[-1].cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = 'attachment reviewed'
    file = tmp_path / 'nested.docx'
    document.save(file)
    audit = run_script('row_audit.py', file, '--source-map', write_json(tmp_path / 'map.json', {'1': 'Reviewed minutes evidence'}))
    assert audit.returncode == 0, audit.stdout
    lint = run_script('lint_minutes.py', file, '--profile', PROFILE)
    assert 'expected an attendance table' not in lint.stdout


def test_cross_year_stale_date_warns_and_blocks_final(tmp_path):
    items = [('1', 'Environment. December 2025 figures remain under review.', '')]
    file = build.build_docx(tmp_path / 'january.docx', items=items)
    r = run_script('lint_minutes.py', file, '--period', 'January',
                   '--month', 'February', '--report-year', '2026', '--final')
    assert r.returncode == 1
    assert 'December 2025 appears in the document' in r.stdout


def test_invalid_hook_input_and_damaged_final_block(tmp_path):
    r = run_script('check_final_docx.py', stdin='not json')
    assert r.returncode == 2
    (tmp_path / 'final.docx').write_bytes(b'corrupted Word file')
    r = run_script('check_final_docx.py', stdin=json.dumps({'cwd': str(tmp_path)}))
    assert r.returncode == 2 and 'cannot validate DOCX' in r.stderr


def test_hook_only_checks_configured_output_directory(tmp_path):
    other = tmp_path / 'unrelated'
    other.mkdir()
    build.build_docx(other / 'final-old.docx', items=[('1', 'Old [TBC]', '')])
    target = tmp_path / 'release'
    target.mkdir()
    build.build_docx(target / 'final-current.docx')
    # A stale final in a subfolder no longer blocks work in the project root.
    assert run_script('check_final_docx.py', stdin=json.dumps({'cwd': str(tmp_path)})).returncode == 0
    env = dict(os.environ, CCC_MINUTES_OUTPUT_DIR=str(other))
    r = run_script('check_final_docx.py', stdin=json.dumps({'cwd': str(tmp_path)}), env=env)
    assert r.returncode == 2


def test_unconfirmed_deck_suggestion_still_blocks(tmp_path):
    # A slide match is a reviewer hint, never a substitute for source-map evidence.
    import shutil
    if not all(shutil.which(cmd) for cmd in ('pdfinfo', 'pdftotext')):
        pytest.skip('Poppler utilities are not installed')
    doc = build.build_docx(tmp_path / 'draft.docx')
    deck = build.build_pdf(tmp_path / 'deck.pdf')
    confirmed = write_json(tmp_path / 'map.json', {
        '1': 'Secretary notes', '3': 'Secretary notes', '4': 'Secretary notes'
    })
    r = run_script('row_audit.py', doc, '--deck', deck, '--source-map', confirmed)
    assert r.returncode == 1, r.stdout
    assert '| Deck slide 1 | **UNCONFIRMED** |' in r.stdout
    assert 'BLOCKER  row 2 has no confirmed source' in r.stdout
    confirmed = write_json(tmp_path / 'map.json', {
        '1': 'Secretary notes', '2': 'Reviewer confirmed deck slide 1',
        '3': 'Secretary notes', '4': 'Secretary notes'
    })
    fixed = run_script('row_audit.py', doc, '--deck', deck, '--source-map', confirmed)
    assert fixed.returncode == 0, fixed.stdout


def test_swap_refuses_path_traversal(tmp_path):
    import importlib.util
    script = Path(__file__).resolve().parent.parent / 'skills/ccc-minutes-builder/scripts/swap_slide_images.py'
    spec = importlib.util.spec_from_file_location('swap_slide_images', script)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    (tmp_path / 'word' / 'media').mkdir(parents=True)
    with pytest.raises(ValueError, match='word/media'):
        module.checked_media_target(tmp_path, '../../unsafe.png')


def test_swap_cleans_temporary_renderings_and_keeps_preview_local(tmp_path, monkeypatch):
    import importlib.util
    script = Path(__file__).resolve().parent.parent / 'skills/ccc-minutes-builder/scripts/swap_slide_images.py'
    spec = importlib.util.spec_from_file_location('swap_slide_images2', script)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    word = tmp_path / 'unpacked' / 'word'
    (word / 'media').mkdir(parents=True)
    (word / '_rels').mkdir()
    Image.new('RGB', (10, 10)).save(word / 'media' / 'image1.png')
    (word / '_rels' / 'document.xml.rels').write_text('''<Relationships>
      <Relationship Id="rId8" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image1.png"/>
    </Relationships>''')
    (word / 'document.xml').write_text('''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
      xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
      xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
      xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">
      <w:body><w:drawing><wp:extent cx="2000" cy="1500"/>
      <a:blip r:embed="rId8"/><a:ext cx="2000" cy="1500"/></w:drawing></w:body></w:document>''')
    render_root = tmp_path / 'renders'
    render_root.mkdir()
    monkeypatch.setattr(module.tempfile, 'tempdir', str(render_root))
    def fake_run(argv, **kwargs):
        Image.new('RGB', (100, 100)).save(str(argv[-1]) + '-1.png')
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(module.subprocess, 'run', fake_run)
    args = SimpleNamespace(unpacked=str(tmp_path / 'unpacked'), rid='rId8',
                           pdf='deck.pdf', page=1, dpi=72, left=0.1,
                           right=0.9, top=0.1, bottom=0.9, preview=None)
    module.cmd_swap(args)
    assert (tmp_path / 'unpacked' / 'preview_rId8.png').is_file()
    assert not list(render_root.iterdir()), 'rendering temp directory must be removed'
    assert Image.open(word / 'media' / 'image1.png').size == (80, 80)
    root = etree.parse(str(word / 'document.xml')).getroot()
    assert root.find('.//' + module.W + 'drawing/' + module.WP + 'extent').get('cy') == '2000'
